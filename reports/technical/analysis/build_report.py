from pathlib import Path
import json,sys
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from report_content import pages
ROOT=Path(__file__).resolve().parents[3];base=ROOT/'reports/technical';out=base/'deliverables'
# Populate the multicorner table directly from machine-readable original metrics.
m=json.loads((base/'evidence/optimized-metrics.json').read_text())
for page in pages:
 if page['title'].startswith('13 '):
  for block in page['blocks']:
   if block[0]=='table':
    for row,key in zip(block[2],['nom_tt_025C_1v80','nom_ss_100C_1v60','nom_ff_n40C_1v95','min_tt_025C_1v80','min_ss_100C_1v60','min_ff_n40C_1v95','max_tt_025C_1v80','max_ss_100C_1v60','max_ff_n40C_1v95']):
     row[1]=f"{m['timing__setup__ws__corner:'+key]:.6f}";row[2]=f"{m['timing__hold__ws__corner:'+key]:.6f}";row[3]=str(m['design__max_slew_violation__count__corner:'+key])
D=Document();sec=D.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11)
sec.top_margin=Inches(.68);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
sec.header_distance=Inches(.25);sec.footer_distance=Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Caption']:
 st=D.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 st._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'Arial')
normal=D.styles['Normal'];normal.font.size=Pt(10.5);normal.paragraph_format.line_spacing=1.10;normal.paragraph_format.space_after=Pt(7)
for name,size in [('Title',25),('Heading 1',17),('Heading 2',12)]:
 st=D.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(7);st.paragraph_format.space_after=Pt(9);st.paragraph_format.keep_with_next=True
D.styles['Caption'].font.size=Pt(9);D.styles['Caption'].font.italic=True;D.styles['Caption'].paragraph_format.space_after=Pt(8)
foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=foot.add_run('TinySystolic  |  ');r.font.size=Pt(8)
f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),'PAGE');foot._p.append(f)
D.core_properties.title='TinySystolic technical report';D.core_properties.subject='Design and pre fabrication verification';D.core_properties.author='Krishna Chemudupati';D.core_properties.keywords='Tiny Tapeout, systolic array, INT4, SKY130, verification'
md=[]
for pi,page in enumerate(pages):
 if pi:D.add_page_break()
 D.add_paragraph(page['title'],'Title' if pi==0 else 'Heading 1');md.append('# '+page['title']+'\n')
 for bi,bl in enumerate(page['blocks']):
  if bl[0]=='p':
   p=D.add_paragraph(bl[1]);md.append(bl[1]+'\n')
   if pi==0 and bi==0:p.style='Subtitle';p.runs[0].font.size=Pt(14)
  elif bl[0]=='h2':D.add_paragraph(bl[1],'Heading 2');md.append('## '+bl[1]+'\n')
  elif bl[0]=='code':
   p=D.add_paragraph();p.paragraph_format.space_after=Pt(8);p.paragraph_format.keep_together=True
   r=p.add_run(bl[1]);r.font.name='Courier New';r.font.size=Pt(9)
   md.append('```text\n'+bl[1]+'\n```\n')
  elif bl[0]=='table':
   _,heads,rows,widths=bl
   t=D.add_table(rows=1,cols=len(heads));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
   widths=widths or [6.8/len(heads)]*len(heads)
   for i,w in enumerate(widths):t.columns[i].width=Inches(w)
   for cell,text in zip(t.rows[0].cells,heads):cell.text=text
   for row in rows:
    cells=t.add_row().cells
    for c,s in zip(cells,row):c.text=s
   for ri,row in enumerate(t.rows):
    trPr=row._tr.get_or_add_trPr();no=OxmlElement('w:cantSplit');trPr.append(no)
    if ri==0:
     repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
    for ci,cell in enumerate(row.cells):
     cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
     pr=cell._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'244B62' if ri==0 else ('F2F5F7' if ri%2==0 else 'FFFFFF'));pr.append(sh)
     mar=OxmlElement('w:tcMar')
     for side,val in [('top',75),('bottom',75),('left',95),('right',95)]:
      el=OxmlElement('w:'+side);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');mar.append(el)
     pr.append(mar)
     bor=OxmlElement('w:tcBorders')
     for side in ['top','left','bottom','right']:
      el=OxmlElement('w:'+side);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bor.append(el)
     pr.append(bor)
     for p in cell.paragraphs:
      p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.03
      if len(heads)>2 and ci>0 and len(cell.text)<24:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
      for r in p.runs:
       r.font.size=Pt(9);r.font.bold=ri==0;r.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
   D.add_paragraph().paragraph_format.space_after=Pt(0)
   md.append('| '+' | '.join(heads)+' |\n| '+' | '.join(['---']*len(heads))+' |\n'+'\n'.join('| '+' | '.join(row)+' |'for row in rows)+'\n')
  elif bl[0]=='figure':
   _,name,cap,width=bl;src=(base/'figures'/name).resolve()
   p=D.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True;p.paragraph_format.space_after=Pt(3)
   shape=p.add_run().add_picture(str(src),width=Inches(width));shape._inline.docPr.set('descr',cap)
   D.add_paragraph(cap,'Caption')
   md.append(f'![{cap}](figures/{name})\n')
  elif bl[0]=='link':
   p=D.add_paragraph();p.paragraph_format.space_after=Pt(3);h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(bl[2],RT.HYPERLINK,is_external=True));r=OxmlElement('w:r');rp=OxmlElement('w:rPr');col=OxmlElement('w:color');col.set(qn('w:val'),'22577A');rp.append(col);r.append(rp);tx=OxmlElement('w:t');tx.text=bl[1];r.append(tx);h.append(r);p._p.append(h);md.append(f'[{bl[1]}]({bl[2]})\n')
for root in [D.styles.element,D.element]:
 for border in root.xpath('.//w:pBdr'):border.getparent().remove(border)
D.save(out/'TinySystolic_Technical_Report.docx')
(base/'TinySystolic_Technical_Report.md').write_text('\n'.join(md))
(base/'report-content.json').write_text(json.dumps(pages,indent=2,ensure_ascii=False)+'\n')
print('Authored',len(pages),'planned pages;',sum(len(x.split()) for x in md),'words including table content')
